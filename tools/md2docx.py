#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Конвертер Markdown -> DOCX для домашнего задания ДЗ 3.

Поддерживает: заголовки #..####, абзацы, маркированные списки (-, вложенные -),
изображения ![alt](путь), горизонтальные линии (---),
нумерованные пункты (как обычные абзацы с сохранением номера), цитаты (>),
таблицы (| ... |), **bold**, *italic*, `code`, ссылки [текст](url),
внутренние ссылки [текст](#anchor), «голые» URL, разрыв строки <br> внутри ячейки.
"""

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

TOKEN = re.compile(
    r"(\*\*.+?\*\*"          # **bold**
    r"|\*[^*\n]+?\*"          # *italic*
    r"|\[[^\]]+\]\([^)\s]+\)"  # [label](url)
    r"|`[^`\n]+`"             # `code`
    r"|https?://[^\s,;)\]]+)"  # bare URL
)

BOLD_ONLY = re.compile(r"^\*\*(.+)\*\*$", re.S)
ITALIC_ONLY = re.compile(r"^\*([^*]+)\*$", re.S)
CODE_ONLY = re.compile(r"^`([^`]+)`$", re.S)
LINK = re.compile(r"^\[([^\]]+)\]\(([^)\s]+)\)$", re.S)

_bid = [1000]


def _new_run(paragraph, text, bold=False, italic=False, code=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if code:
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        rpr = run._element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(a), "Consolas")
    return run


def add_hyperlink(paragraph, url, text):
    """Внешняя ссылка (url) или внутренняя (#anchor)."""
    if url.startswith("#"):
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("w:anchor"), url[1:])
    else:
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hl.append(run)
    paragraph._p.append(hl)


def add_bookmark(paragraph, name):
    _bid[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bid[0]))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bid[0]))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def emit(paragraph, text, bold=False, italic=False, mono=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        m = BOLD_ONLY.match(part)
        if m:
            emit(paragraph, m.group(1), True, italic, mono)
            continue
        m = ITALIC_ONLY.match(part)
        if m:
            emit(paragraph, m.group(1), bold, True, mono)
            continue
        m = CODE_ONLY.match(part)
        if m:
            _new_run(paragraph, m.group(1), bold, italic, True)
            continue
        m = LINK.match(part)
        if m:
            add_hyperlink(paragraph, m.group(2), m.group(1))
            continue
        if part.startswith("http"):
            url = part
            tail = ""
            while url and url[-1] in ".,:;!?":
                tail = url[-1] + tail
                url = url[:-1]
            add_hyperlink(paragraph, url, url)
            if tail:
                _new_run(paragraph, tail, bold, italic, mono)
            continue
        _new_run(paragraph, part, bold, italic, mono)


def fill_cell(cell, md):
    cell.text = ""
    for i, part in enumerate(md.split("<br>")):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        emit(p, part.strip())


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)


def shade(cell, color="EDF2F7"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            md = row[j] if j < len(row) else ""
            fill_cell(t.cell(i, j), md)
            if i == 0:
                shade(t.cell(i, j))
                for p in t.cell(i, j).paragraphs:
                    for run in p.runs:
                        run.bold = True
    set_repeat_header(t.rows[0])
    # узкая колонка "№" в сводной таблице
    if rows and rows[0][0].strip() in ("№", "N"):
        widths = [Cm(1.1), Cm(4.6), Cm(11.2)]
        for r in t.rows:
            for j, w in enumerate(widths):
                if j < ncols:
                    r.cells[j].width = w
    doc.add_paragraph()


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip()
        if re.match(r"^\|[\s:|-]+\|$", raw):
            i += 1
            continue
        cells = [c.strip() for c in raw.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def build(md_path, docx_path):
    with open(md_path, encoding="utf-8") as fh:
        lines = fh.read().split("\n")

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), "Calibri")

    for s in doc.sections:
        s.left_margin = Cm(2)
        s.right_margin = Cm(1.5)
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("<a ") or stripped.startswith("<!--"):
            if "Список источников" in stripped:
                pass
            i += 1
            continue

        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if m:
            img_path = m.group(2)
            if not os.path.isabs(img_path):
                img_path = os.path.join(os.path.dirname(os.path.abspath(md_path)), img_path)
            if os.path.exists(img_path):
                doc.add_picture(img_path, width=Cm(15.5))
                doc.paragraphs[-1].alignment = 1
                if m.group(1):
                    cap = doc.add_paragraph()
                    cap.alignment = 1
                    r = cap.add_run(m.group(1))
                    r.italic = True
                    r.font.size = Pt(9)
            else:
                p = doc.add_paragraph()
                emit(p, f"[изображение не найдено: {img_path}]", italic=True)
            i += 1
            continue

        if re.match(r"^-{3,}$", stripped):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "BBBBBB")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=level)
            emit(h, m.group(2).strip())
            if m.group(2).strip().startswith("11. Список источников"):
                add_bookmark(h, "sources")
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(4)
            emit(p, stripped.lstrip("> ").strip(), italic=True)
            i += 1
            continue

        m = re.match(r"^\s{2,}-\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet 2")
            emit(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^-\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            emit(p, m.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        emit(p, stripped)
        i += 1

    doc.save(docx_path)
    print("saved:", docx_path)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
